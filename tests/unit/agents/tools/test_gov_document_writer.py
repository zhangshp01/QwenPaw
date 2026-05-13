# -*- coding: utf-8 -*-
from __future__ import annotations

import asyncio
import json
import tempfile
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest

from qwenpaw.agents.tools.gov_document_writer import (
    _parse_document_review_response,
    doc_reviewer,
    gov_document_layout,
    gov_document_writer,
)
from qwenpaw.config.context import (
    clear_agent_tool_step_for_running_task,
    set_agent_tool_step_for_running_task,
    set_current_workspace_dir,
)

_DOC_REVIEW_PATCH = "qwenpaw.agents.tools.gov_document_writer._run_document_review_llm"


def _json_payload(r):
    block = r.content[0]
    assert isinstance(block, dict)
    assert block.get("type") == "json"
    return block["json"]


def _docx_combined_text(save_path: str) -> str:
    from docx import Document as DocxDocument

    d = DocxDocument(save_path)
    return "\n".join(para.text for para in d.paragraphs)


def test_doc_reviewer_requires_input():
    async def _run():
        r = await doc_reviewer(content="", file_path="", path="")
        p = _json_payload(r)
        assert p["skillName"] == "doc_reviewer"
        assert p["sourceState"] == "error"
        assert p["normalizedResult"] is None
        assert p["displayText"] == "文档审核失败"
        detail = p.get("errorDetail") or ""
        assert "content" in detail and ("file_path" in detail or "path" in detail)
        assert p.get("resultList") is None

    asyncio.run(_run())
    async def _run():
        revised = "para1-fixed.\npara2-fixed."
        fake = AsyncMock(return_value=(revised, []))
        with patch(_DOC_REVIEW_PATCH, fake):
            r = await doc_reviewer(content="para1.\npara2.")
        p = _json_payload(r)
        assert p["skillName"] == "doc_reviewer"
        assert p["sourceState"] == "model_success"
        assert p["displayText"] == "已完成：文档审核"
        assert p["stepIndex"] == 1
        nr = p["normalizedResult"]
        assert nr["source"] == "model_success"
        assert nr["document"] == revised
        assert p["resultList"] == []
        fake.assert_awaited_once()
        assert "para1" in (fake.await_args.args[0] or "")

    asyncio.run(_run())


def test_doc_reviewer_reads_file():
    async def _run():
        fake = AsyncMock(return_value=("hello-revised", []))
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            p = root / "draft.md"
            p.write_text("# T\n\nhello", encoding="utf-8")
            with patch(_DOC_REVIEW_PATCH, fake):
                r = await doc_reviewer(file_path="draft.md")
            payload = _json_payload(r)
            assert payload["skillName"] == "doc_reviewer"
            assert payload["displayText"] == "已完成：文档审核"
            assert payload["normalizedResult"]["document"] == "hello-revised"
            fake.assert_awaited_once()
            assert "hello" in (fake.await_args.args[0] or "")

    asyncio.run(_run())


def test_doc_reviewer_step_index_from_task_mapping():
    async def _run():
        fake = AsyncMock(return_value=("final-body", []))
        set_agent_tool_step_for_running_task(2)
        try:
            with patch(_DOC_REVIEW_PATCH, fake):
                r = await doc_reviewer(content="x")
            p = _json_payload(r)
            assert p["stepIndex"] == 2
        finally:
            clear_agent_tool_step_for_running_task()

    asyncio.run(_run())


def test_doc_reviewer_llm_failure_returns_error():
    async def _run():
        fake = AsyncMock(side_effect=RuntimeError("model-unavailable"))
        with patch(_DOC_REVIEW_PATCH, fake):
            r = await doc_reviewer(content="one line.")
        p = _json_payload(r)
        assert p["sourceState"] == "error"
        assert p["normalizedResult"] is None
        assert "model-unavailable" in (p.get("errorDetail") or "")
        assert p.get("resultList") is None

    asyncio.run(_run())


def test_doc_reviewer_result_list_from_llm_tuple():
    async def _run():
        items = [
            {
                "reason": "错别字",
                "offsets": 25,
                "errorType": "错别字错误",
                "errorWord": "法展",
                "contextOffset": 10,
                "context": "随着大数据技术的迅速法展，",
                "rightWord": "发展",
            },
        ]
        fake = AsyncMock(return_value=("定稿正文", items))
        with patch(_DOC_REVIEW_PATCH, fake):
            r = await doc_reviewer(content="x")
        p = _json_payload(r)
        assert p["normalizedResult"]["document"] == "定稿正文"
        assert len(p["resultList"]) == 1
        row = p["resultList"][0]
        assert row["errorWord"] == "法展"
        assert row["rightWord"] == "发展"
        assert row["offsets"] == 25
        assert row["contextOffset"] == 10

    asyncio.run(_run())


def test_parse_document_review_response_json():
    raw = json.dumps(
        {
            "document": "后文",
            "resultList": [
                {
                    "reason": "r",
                    "offsets": 1,
                    "errorType": "t",
                    "errorWord": "错",
                    "contextOffset": 0,
                    "context": "上下文",
                    "rightWord": "对",
                },
            ],
        },
        ensure_ascii=False,
    )
    doc, items = _parse_document_review_response(raw)
    assert doc == "后文"
    assert len(items) == 1
    assert items[0]["errorWord"] == "错"


def test_parse_document_review_response_plain_text_legacy():
    doc, items = _parse_document_review_response("  仅正文无JSON  ")
    assert doc == "仅正文无JSON"
    assert items == []


def test_parse_document_review_response_invalid_document_key():
    with pytest.raises(RuntimeError, match="document"):
        _parse_document_review_response('{"document": ""}')


def test_gov_document_writer_requires_content():
    async def _run():
        r = await gov_document_writer(content="")
        p = _json_payload(r)
        assert p["normalizedResult"] is None
        assert p["sourceState"] == "error"
        assert p["retryable"] is True
        assert p["savePath"] is None
        assert p["skillName"] == "gov-document-writer"
        assert "content" in (p.get("errorDetail") or "").lower()

    asyncio.run(_run())


def test_gov_document_writer_writes_file():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            r = await gov_document_writer(
                title="Holiday Notice",
                content="Body line one.",
                type="notice",
            )
            p = _json_payload(r)
            assert p["skillName"] == "gov-document-writer"
            assert p["sourceState"] == "model_success"
            assert p["errorDetail"] is None
            assert p["retryable"] is True
            assert p["stepIndex"] == 1
            sp = (p.get("savePath") or "").replace("\\", "/")
            assert "govdocs" in sp
            assert sp.lower().endswith(".docx")
            nr = p["normalizedResult"]
            assert nr["source"] == "model_success"
            assert "Body line one." in nr["document"]
            assert "**类型**" not in nr["document"]
            assert "notice" not in nr["document"]
            docx_files = list((root / "govdocs").glob("Holiday Notice-*.docx"))
            assert len(docx_files) == 1
            assert Path(p["savePath"]).resolve() == docx_files[0].resolve()
            from docx import Document as DocxDocument

            d = DocxDocument(str(docx_files[0]))
            combined = "\n".join(para.text for para in d.paragraphs)
            assert "Holiday Notice" in combined
            assert "Body line one." in combined

    asyncio.run(_run())


def test_gov_document_layout_from_file_path_only():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            (root / "draft.md").write_text("从文件读取的正文。", encoding="utf-8")
            r = await gov_document_layout(
                title="文件标题",
                file_path="draft.md",
            )
            p = _json_payload(r)
            assert p["skillName"] == "gov_document_layout"
            assert p["sourceState"] == "model_success"
            assert "normalizedResult" not in p
            assert p["resultList"] == {"total": 0, "list": []}
            combined = _docx_combined_text(p["savePath"])
            assert "从文件读取的正文" in combined

    asyncio.run(_run())


def test_gov_document_layout_main_text_camel_param():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            r = await gov_document_layout(
                title="Camel",
                content="",
                mainText="camelCase 字段正文。",
            )
            p = _json_payload(r)
            assert p["sourceState"] == "model_success"
            assert "normalizedResult" not in p
            assert p["resultList"]["total"] == 0
            combined = _docx_combined_text(p["savePath"])
            assert "camelCase" in combined

    asyncio.run(_run())


def test_gov_document_layout_data_json_string():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            r = await gov_document_layout(
                title="JSON",
                data='{"content": "嵌套 JSON 里的正文。"}',
            )
            p = _json_payload(r)
            assert p["sourceState"] == "model_success"
            assert "normalizedResult" not in p
            combined = _docx_combined_text(p["savePath"])
            assert "嵌套 JSON" in combined

    asyncio.run(_run())


def test_gov_document_layout_extra_unknown_long_key():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            r = await gov_document_layout(
                title="T",
                **{"segmentContent": "这是唯一的长字段内容。"},
            )
            p = _json_payload(r)
            assert p["sourceState"] == "model_success"
            assert "normalizedResult" not in p
            combined = _docx_combined_text(p["savePath"])
            assert "长字段" in combined

    asyncio.run(_run())


def test_gov_document_layout_requires_content():
    async def _run():
        r = await gov_document_layout(
            content="",
            keyword="",
            text="",
            document="",
            body="",
        )
        p = _json_payload(r)
        assert p["normalizedResult"] is None
        assert p["sourceState"] == "error"
        assert p["skillName"] == "gov_document_layout"

    asyncio.run(_run())


def test_gov_document_layout_body_from_keyword_when_content_empty():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            r = await gov_document_layout(
                title="From Keyword",
                content="",
                keyword="仅通过 keyword 传入的正文段落。",
            )
            p = _json_payload(r)
            assert p["skillName"] == "gov_document_layout"
            assert p["sourceState"] == "model_success"
            assert "normalizedResult" not in p
            assert "仅通过 keyword 传入的正文段落" in _docx_combined_text(p["savePath"])

    asyncio.run(_run())


def test_gov_document_layout_writes_file():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            r = await gov_document_layout(
                title="Layout Test",
                content="Paragraph for layout.",
            )
            p = _json_payload(r)
            assert p["skillName"] == "gov_document_layout"
            assert p["sourceState"] == "model_success"
            assert p["displayText"] == "已完成：公文排版"
            assert "normalizedResult" not in p
            assert p["resultList"] == {"total": 0, "list": []}
            sp = (p.get("savePath") or "").replace("\\", "/")
            assert "govdocs" in sp
            assert sp.lower().endswith(".docx")

    asyncio.run(_run())


def test_gov_document_layout_accepts_keyword_kwarg():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            r = await gov_document_layout(
                title="K",
                content="body",
                keyword="ui-metadata",
            )
            p = _json_payload(r)
            assert p["skillName"] == "gov_document_layout"
            assert p["sourceState"] == "model_success"
            assert "normalizedResult" not in p

    asyncio.run(_run())


def test_gov_document_layout_result_list_from_templates_json():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            set_current_workspace_dir(root)
            templates = json.dumps(
                [
                    {
                        "id": "7eac05530af3445fa3dccf343e8c4e36",
                        "templateTitle": "市局党委上行文",
                        "industryTag": "昆明",
                        "share": True,
                        "templateType": "红头",
                        "layoutContent": "",
                        "count": 10,
                    },
                    {
                        "id": "a196162515d04870bd4f3b38dc548106",
                        "templateTitle": "市局党委平行、下行文",
                        "documentType": "kmdangweixiaxingwen",
                        "industryTag": "昆明",
                        "share": True,
                        "templateType": "红头",
                        "layoutContent": "",
                        "count": 9,
                    },
                ],
                ensure_ascii=False,
            )
            r = await gov_document_layout(
                title="示例标题",
                content="这是一段示例正文内容。",
                templates=templates,
            )
            p = _json_payload(r)
            assert "normalizedResult" not in p
            rl = p["resultList"]
            assert rl["total"] == 2
            assert len(rl["list"]) == 2
            assert rl["list"][0]["templateTitle"] == "市局党委上行文"
            assert rl["list"][1]["documentType"] == "kmdangweixiaxingwen"

    asyncio.run(_run())


def test_gov_document_layout_query_only_mocked_fetch():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            set_current_workspace_dir(Path(td))
            with patch(
                "qwenpaw.agents.tools.gov_document_writer.fetch_layout_template_result_list",
                return_value={"total": 1, "list": [{"id": "1"}]},
            ) as mock_fetch:
                r = await gov_document_layout(template_title="上行文")
            mock_fetch.assert_called_once()
            p = _json_payload(r)
            assert p["savePath"] is None
            assert p["sourceState"] == "model_success"
            assert p["resultList"]["total"] == 1

    asyncio.run(_run())


def test_gov_document_layout_content_and_template_title_uses_fetched_result_list():
    async def _run():
        with tempfile.TemporaryDirectory() as td:
            set_current_workspace_dir(Path(td))
            with patch(
                "qwenpaw.agents.tools.gov_document_writer.fetch_layout_template_result_list",
                return_value={"total": 3, "list": [{"id": "a"}]},
            ):
                r = await gov_document_layout(
                    template_title="党委",
                    content="正文",
                    title="T",
                )
            p = _json_payload(r)
            assert p["savePath"]
            assert p["resultList"]["total"] == 3
            assert len(p["resultList"]["list"]) == 1

    asyncio.run(_run())


def test_gov_document_writer_step_index_from_task_mapping():
    async def _run():
        set_agent_tool_step_for_running_task(3)
        try:
            r = await gov_document_writer(content="")
            p = _json_payload(r)
            assert p["stepIndex"] == 3
        finally:
            clear_agent_tool_step_for_running_task()

    asyncio.run(_run())
