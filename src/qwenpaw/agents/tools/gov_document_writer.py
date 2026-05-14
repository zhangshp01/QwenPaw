# -*- coding: utf-8 -*-
"""Formal / gov-style document draft (``gov_document_writer``), layout
(``gov_document_layout``), and review (``doc_reviewer``) helpers for AgentLoop.

``gov_document_writer`` returns one ``{"type": "json", "json": {...}}`` block with
``normalizedResult.document`` (title + body plain text) and **does not** write
files or emit ``savePath``.

``gov_document_layout`` only recommends templates from HaiRuo using non-empty
``template_title`` (the 公文标题); it does **not** write ``.docx`` files. Success
responses include a non-null ``savePath`` string as the **pending** output path
(same naming rule as before: ``govdocs/{stem}-{timestamp}.docx`` unless
``savePath`` / ``save_path`` hints an absolute/relative path).

``doc_reviewer`` returns ``normalizedResult`` plus ``resultList``.  Layout
payloads may include ``recommended`` only (see ``hairuo_gov_layout``).
"""

from __future__ import annotations

import asyncio
import json
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from agentscope.tool import ToolResponse

from ...config.context import (
    get_agent_tool_step_for_running_task,
    get_current_recent_max_bytes,
)
from .file_io import _resolve_file_path, read_file
from .hairuo_gov_layout import GovLayoutHttpError, fetch_layout_template_result_list
from .utils import DEFAULT_MAX_BYTES, truncate_text_output

_SKILL_NAME_EN = "gov-document-writer"
_SKILL_NAME_LAYOUT_EN = "gov_document_layout"
_SKILL_NAME_DOC_REVIEWER_EN = "doc_reviewer"
_DISPLAY_OK_ZH = "已完成：公文写作"
_DISPLAY_ERR_ZH = "公文写作失败"
_DISPLAY_OK_LAYOUT_ZH = "已完成：公文排版"
_DISPLAY_ERR_LAYOUT_ZH = "公文排版失败"
_DISPLAY_OK_DOC_REVIEW_ZH = "已完成：文档审核"
_DISPLAY_ERR_DOC_REVIEW_ZH = "文档审核失败"
_SOURCE_OK = "model_success"
_SOURCE_ERR = "error"
_FALLBACK_STEP_INDEX = 1

logger = logging.getLogger(__name__)

_DOC_REVIEW_LLM_TIMEOUT = 600.0
_DOC_REVIEW_MAX_USER_CHARS = 48_000
_DOC_REVIEW_SYSTEM_ZH = (
    "你是中文文档校对与润色助手。用户将提供一段待审正文（可能含错别字或表述问题）。\n"
    "请**只输出一个 JSON 对象**（不要 Markdown 代码围栏、不要输出任何 JSON 以外的说明文字）。\n"
    "JSON 顶层字段必须为：\n"
    '  "document": 字符串，审核修改后的完整正文（可直接替换原文的定稿）；\n'
    '  "resultList": 数组，列出相对用户原文你实际改动过的要点（无改动则为 []）。\n'
    "resultList 中每一项为对象，字段与含义如下（均为与用户所给**原文**对照）：\n"
    '  "reason": 判定与修改依据（字符串）；\n'
    '  "offsets": 问题片段在原文中的起始字符偏移（非负整数，按 Unicode 字符索引从 0 起）；\n'
    '  "errorType": 问题类型，如「错别字错误」「标点错误」「表述问题」等；\n'
    '  "errorWord": 原文中的问题片段；\n'
    '  "contextOffset": 用于展示的上下文在原文中的起始偏移（非负整数）；\n'
    '  "context": 含问题的短上下文（字符串）；\n'
    '  "rightWord": 建议替换为的正确写法（删除类可为空字符串）。\n'
    "正文要求：\n"
    "1）修正错别字、明显用词与标点错误；可在不改变原意的前提下轻微润色；\n"
    "2）尽量保留原文段落换行与层次，不要擅自改成分析报告或条目清单；\n"
    "3）resultList 应与 document 中的修改一致；若无法给出精确 offsets，可填合理近似整数。"
)


def _content_block_text(block: Any) -> str | None:
    if isinstance(block, dict):
        return block.get("text")
    return getattr(block, "text", None)


def _safe_stem(name: str, max_len: int = 80) -> str:
    name = (name or "").strip() or "document"
    name = Path(name).name
    cleaned = re.sub(r'[<>:"/\\|?*\n\r\t]', "_", name)
    cleaned = cleaned.strip(" .") or "document"
    return cleaned[:max_len]


def _docx_basename(heading: str) -> str:
    """``{公文标题}-{时间戳}.docx`` (timestamp: local ``%Y%m%d%H%M%S`` + microseconds)."""
    ts = datetime.now().strftime("%Y%m%d%H%M%S%f")
    return f"{_safe_stem(heading)}-{ts}.docx"


def _resolve_gov_docx_output_path(save_hint: str, heading: str) -> Path:
    """Resolve where to write the ``.docx`` (optional ``savePath`` / ``save_path``)."""
    hint = (save_hint or "").strip()
    if not hint:
        rel = f"govdocs/{_docx_basename(heading)}"
        return Path(_resolve_file_path(rel))
    out = Path(_resolve_file_path(hint))
    if out.suffix.lower() != ".docx":
        out = out.with_suffix(".docx")
    return out


def _normalize_layout_template_list_payload(
    raw: Any,
    *,
    templates_json: str = "",
) -> dict[str, Any]:
    """Return ``{"templates": [...]}`` for ``gov_document_layout`` payloads.

    Accepts legacy shapes ``{"list", "total"}`` when callers pass explicit ``resultList``.
    """
    empty: dict[str, Any] = {"templates": []}

    def from_list(lst: list[Any]) -> dict[str, Any]:
        items = [x for x in lst if isinstance(x, dict)]
        return {"templates": items}

    if isinstance(raw, str):
        s = raw.strip()
        if s:
            try:
                raw = json.loads(s)
            except json.JSONDecodeError:
                return empty
        else:
            raw = None

    if isinstance(raw, dict):
        t_new = raw.get("templates")
        if isinstance(t_new, list):
            items = [x for x in t_new if isinstance(x, dict)]
            return {"templates": items}
        lst = raw.get("list")
        if isinstance(lst, list):
            items = [x for x in lst if isinstance(x, dict)]
            return {"templates": items}
        return empty
    if isinstance(raw, list):
        return from_list(raw)

    tj = (templates_json or "").strip()
    if tj:
        try:
            arr = json.loads(tj)
            if isinstance(arr, list):
                return from_list(arr)
        except json.JSONDecodeError:
            return empty
    return empty


def _layout_result_list_from_tool_locals(_locals: dict[str, Any]) -> dict[str, Any]:
    """Resolve ``resultList`` / ``templates`` / ``**extra`` for layout tool."""
    raw: Any = _locals.get("resultList")
    templates = str(_locals.get("templates") or "")
    extra = _locals.get("extra")
    if (raw is None or (isinstance(raw, str) and not raw.strip())) and isinstance(
        extra,
        dict,
    ):
        if extra.get("resultList") is not None:
            raw = extra.get("resultList")
        if not templates.strip():
            ex_t = extra.get("templates")
            if isinstance(ex_t, str):
                templates = ex_t
    out = _normalize_layout_template_list_payload(raw, templates_json=templates)
    if not (out.get("templates") or []) and isinstance(extra, dict):
        ex_t2 = extra.get("templates")
        if isinstance(ex_t2, str) and ex_t2.strip():
            out = _normalize_layout_template_list_payload(
                None,
                templates_json=ex_t2,
            )
    return out


_LAYOUT_RESULT_CLIENT_DROP: frozenset[str] = frozenset({"templates", "list", "total"})


def _layout_result_list_for_client(rl: dict[str, Any]) -> dict[str, Any] | None:
    """Omit ``templates`` / legacy ``list`` / ``total`` from layout tool JSON."""
    if not isinstance(rl, dict):
        return None
    out = {k: v for k, v in rl.items() if k not in _LAYOUT_RESULT_CLIENT_DROP}
    return out if out else None


def _write_gov_docx(
    path: Path,
    heading: str,
    body: str,
    *,
    type_label: str = "",
) -> None:
    """Blocking: build a minimal .docx (requires ``python-docx``).

    Optional ``type_label`` (e.g. 中文「通知」) is shown as one line
    ``类型：…``; omitted when empty so machine tokens like ``notice`` are not
    injected into formal text.
    """
    try:
        from docx import Document  # type: ignore[import-untyped]
    except ImportError as exc:
        raise RuntimeError(
            "缺少依赖 python-docx，无法生成 .docx。请在运行环境中安装：pip install python-docx",
        ) from exc

    doc = Document()
    doc.add_heading(heading, level=0)
    label = (type_label or "").strip()
    if label:
        doc.add_paragraph(f"类型：{label}")
    for line in body.split("\n"):
        doc.add_paragraph(line.rstrip("\r"))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _writer_step_index() -> int:
    """1-based tool order in the current agent ``reply`` (toolkit middleware)."""
    step = get_agent_tool_step_for_running_task()
    return int(step) if step is not None else _FALLBACK_STEP_INDEX


_GOV_INLINE_BODY_KEYS: tuple[str, ...] = (
    "content",
    "text",
    "document",
    "body",
    "markdown",
    "full_text",
    "plain_text",
    "article",
    "source_text",
    "message",
    "draft",
    "query",
    "prompt",
    "input",
    "mainText",
    "main_text",
    "plainText",
    "fullText",
    "userContent",
    "sourceText",
    "keyword",
)

_GOV_JSON_NEST_KEYS: tuple[str, ...] = (
    "params",
    "parameters",
    "payload",
    "data",
    "args",
    "arguments",
    "extra",
)

_ALL_BODY_VALUE_KEYS: tuple[str, ...] = _GOV_INLINE_BODY_KEYS + ("data",)

# Keys whose string values must not be used as a loose ``**extra`` body guess.
_GOV_EXTRA_VALUE_SKIP_KEYS: frozenset[str] = frozenset(
    {
        "title",
        "file_name",
        "document_type",
        "type",
        "file_path",
        "path",
        "save_path",
        "savePath",
        "resultList",
        "templates",
        "template_title",
    },
)


def _longest_extra_string_body(extra: dict[str, Any], *, min_len: int) -> str:
    """Last resort: longest non-skipped string in ``**extra`` (unknown field names)."""
    best = ""
    for key, val in extra.items():
        if key in _GOV_EXTRA_VALUE_SKIP_KEYS or str(key).startswith("_"):
            continue
        if not isinstance(val, str):
            continue
        text = val.strip()
        if len(text) < min_len:
            continue
        if len(text) > len(best):
            best = text
    return best


def _body_from_maybe_json(text: str) -> str:
    raw = (text or "").strip()
    if not raw or not raw.startswith("{"):
        return ""
    try:
        obj: Any = json.loads(raw)
    except json.JSONDecodeError:
        return ""
    if isinstance(obj, dict):
        return _extract_body_from_mapping(obj)
    return ""


def _extract_body_from_mapping(obj: dict[str, Any], depth: int = 0) -> str:
    """Pull first usable body string from nested AgentLoop-style JSON."""
    if depth > 8:
        return ""
    for key in _GOV_INLINE_BODY_KEYS:
        val = obj.get(key)
        if isinstance(val, str) and val.strip():
            return val.strip()
    for nest in _GOV_JSON_NEST_KEYS:
        inner = obj.get(nest)
        if isinstance(inner, dict):
            got = _extract_body_from_mapping(inner, depth + 1)
            if got:
                return got
        if isinstance(inner, str) and (s := inner.strip()):
            if s.startswith("{"):
                got = _body_from_maybe_json(s)
                if got:
                    return got
    return ""


def _merge_gov_inline_body(
    flat: dict[str, Any],
    extra: dict[str, Any] | None = None,
) -> str:
    """Prefer explicit ``flat`` fields, then ``**extra``, then nested JSON."""

    def _scan_strings(block: dict[str, Any]) -> str:
        for key in _GOV_INLINE_BODY_KEYS:
            val = block.get(key)
            if isinstance(val, str) and val.strip():
                return val.strip()
        return ""

    def _scan_data(block: dict[str, Any]) -> str:
        dv = block.get("data")
        if isinstance(dv, dict):
            got = _extract_body_from_mapping(dv)
            if got:
                return got
        if isinstance(dv, str) and (s := dv.strip()):
            got = _body_from_maybe_json(s)
            if got:
                return got
            return s
        return ""

    got = _scan_strings(flat)
    if got:
        return got
    got = _scan_data(flat)
    if got:
        return got
    if extra:
        got = _scan_strings(extra)
        if got:
            return got
        got = _scan_data(extra)
        if got:
            return got
        got = _extract_body_from_mapping(extra)
        if got:
            return got
    got = _extract_body_from_mapping(flat)
    if got:
        return got
    if extra:
        loose = _longest_extra_string_body(extra, min_len=1)
        if loose:
            return loose
    return ""


def _merge_body_from_tool_locals(_locals: dict[str, Any]) -> str:
    """Build merge dict from a gov tool's ``locals()`` (expects ``extra`` dict)."""
    extra_raw = _locals.get("extra")
    extra = extra_raw if isinstance(extra_raw, dict) else {}
    flat = {k: _locals.get(k, "") for k in _ALL_BODY_VALUE_KEYS}
    return _merge_gov_inline_body(flat, extra)


def _gov_doc_tool_json(root: dict[str, Any]) -> ToolResponse:
    return ToolResponse(content=[{"type": "json", "json": root}])


def _gov_layout_ok_json_recommend_only(
    result_list: dict[str, Any],
    *,
    pending_save_path: str,
) -> ToolResponse:
    """Template recommendation only: no ``.docx`` on disk; ``savePath`` is pending."""
    root: dict[str, Any] = {
        "skillName": _SKILL_NAME_LAYOUT_EN,
        "stepIndex": _writer_step_index(),
        "displayText": _DISPLAY_OK_LAYOUT_ZH,
        "retryable": True,
        "sourceState": _SOURCE_OK,
        "errorDetail": None,
        "savePath": pending_save_path,
    }
    client_rl = _layout_result_list_for_client(result_list)
    if client_rl is not None:
        root["resultList"] = client_rl
    return _gov_doc_tool_json(root)


def _gov_layout_heading_for_pending_path(
    template_title: str,
    title: str,
    file_name: str,
) -> str:
    """Stem for pending ``.docx`` path (same precedence as layout docx naming)."""
    heading = (title or "").strip()
    if not heading and (file_name or "").strip():
        heading = Path(file_name.strip()).stem
    if not heading:
        heading = (template_title or "").strip() or "公文排版稿"
    return heading


def _gov_layout_pending_save_absolute_path(
    *,
    template_title: str,
    title: str,
    file_name: str,
    save_path: str,
    savePath: str,
) -> str:
    heading = _gov_layout_heading_for_pending_path(template_title, title, file_name)
    hint = (savePath or save_path or "").strip()
    resolved = _resolve_gov_docx_output_path(hint, heading)
    return str(resolved.resolve())


def _gov_doc_tool_error(
    skill_name_en: str,
    display_err_zh: str,
    detail: str,
) -> ToolResponse:
    detail = (detail or "").strip() or "unknown"
    root: dict[str, Any] = {
        "skillName": skill_name_en,
        "stepIndex": _writer_step_index(),
        "displayText": display_err_zh,
        "retryable": True,
        "sourceState": _SOURCE_ERR,
        "errorDetail": detail,
        "normalizedResult": None,
    }
    if skill_name_en != _SKILL_NAME_EN:
        root["savePath"] = None
    return _gov_doc_tool_json(root)


def _doc_reviewer_json(root: dict[str, Any]) -> ToolResponse:
    return ToolResponse(content=[{"type": "json", "json": root}])


def _doc_reviewer_error(detail: str) -> ToolResponse:
    detail = (detail or "").strip() or "unknown"
    return _doc_reviewer_json(
        {
            "skillName": _SKILL_NAME_DOC_REVIEWER_EN,
            "stepIndex": _writer_step_index(),
            "displayText": _DISPLAY_ERR_DOC_REVIEW_ZH,
            "retryable": True,
            "sourceState": _SOURCE_ERR,
            "errorDetail": detail,
            "normalizedResult": None,
            "resultList": None,
        },
    )


def _doc_reviewer_ok(document: str, result_list: list[dict[str, Any]]) -> ToolResponse:
    return _doc_reviewer_json(
        {
            "skillName": _SKILL_NAME_DOC_REVIEWER_EN,
            "stepIndex": _writer_step_index(),
            "displayText": _DISPLAY_OK_DOC_REVIEW_ZH,
            "retryable": True,
            "sourceState": _SOURCE_OK,
            "errorDetail": None,
            "normalizedResult": {
                "document": document,
                "source": _SOURCE_OK,
            },
            "resultList": result_list,
        },
    )


_GOV_WRITER_SKILL_NAMES = frozenset({"gov-document-writer", "gov_document_writer"})


def _unwrap_gov_writer_prior_document(raw: Any) -> str:
    """Return ``normalizedResult.document`` from a prior ``gov-document-writer`` payload."""
    if raw is None:
        return ""
    if isinstance(raw, str):
        s = raw.strip()
        if not s:
            return ""
        if s.startswith(("{", "[")):
            try:
                return _unwrap_gov_writer_prior_document(json.loads(s))
            except json.JSONDecodeError:
                return ""
        return ""
    if isinstance(raw, dict):
        if raw.get("type") == "json" and isinstance(raw.get("json"), dict):
            return _unwrap_gov_writer_prior_document(raw["json"])
        sk_raw = raw.get("skillName") or raw.get("skill_name")
        sk = str(sk_raw).strip() if sk_raw is not None else ""
        if sk == "doc_reviewer":
            return ""
        if sk and sk not in _GOV_WRITER_SKILL_NAMES:
            return ""
        nr = raw.get("normalizedResult")
        if isinstance(nr, dict):
            doc = nr.get("document")
            if isinstance(doc, str) and doc.strip():
                return doc.strip()
        return ""
    if isinstance(raw, list):
        for item in raw:
            got = _unwrap_gov_writer_prior_document(item)
            if got:
                return got
    return ""


def _strip_outer_code_fence(text: str) -> str:
    """If the model wrapped the body in a Markdown fence, return inner text."""
    t = (text or "").strip()
    if not t.startswith("```"):
        return t
    lines = t.split("\n")
    if not lines:
        return t
    if lines[0].startswith("```"):
        lines = lines[1:]
    if lines and lines[-1].strip() == "```":
        lines = lines[:-1]
    return "\n".join(lines).strip()


def _coerce_offset(val: Any) -> int | None:
    if isinstance(val, bool):
        return None
    if isinstance(val, int) and val >= 0:
        return val
    if isinstance(val, float) and val >= 0 and val.is_integer():
        return int(val)
    if isinstance(val, str) and val.strip():
        try:
            n = int(val.strip())
            return n if n >= 0 else None
        except ValueError:
            return None
    return None


def _normalize_review_item(raw: dict[str, Any]) -> dict[str, Any]:
    """Shape one resultList entry for stable JSON (frontend / 工作区技能)."""
    off = _coerce_offset(raw.get("offsets"))
    ctx_off = _coerce_offset(raw.get("contextOffset"))
    return {
        "reason": str(raw.get("reason") or "").strip(),
        "offsets": off,
        "errorType": str(raw.get("errorType") or "").strip(),
        "errorWord": str(raw.get("errorWord") or "").strip(),
        "contextOffset": ctx_off,
        "context": str(raw.get("context") or "").strip(),
        "rightWord": str(raw.get("rightWord") or "").strip(),
    }


def _parse_document_review_response(raw: str) -> tuple[str, list[dict[str, Any]]]:
    """Parse model output: JSON with document + resultList, else whole text + []."""
    text = _strip_outer_code_fence(raw or "").strip()
    if not text:
        raise RuntimeError("模型未返回内容。")

    data: Any = None
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        brace = text.find("{")
        if brace >= 0:
            try:
                data, _ = json.JSONDecoder().raw_decode(text[brace:])
            except json.JSONDecodeError:
                data = None

    if isinstance(data, dict):
        doc = data.get("document")
        if isinstance(doc, str) and doc.strip():
            items_raw = data.get("resultList")
            out_list: list[dict[str, Any]] = []
            if isinstance(items_raw, list):
                for it in items_raw:
                    if isinstance(it, dict):
                        out_list.append(_normalize_review_item(it))
            return doc.strip(), out_list
        raise RuntimeError("模型返回的 JSON 缺少非空的 document 字符串。")

    if data is not None:
        raise RuntimeError(
            "模型返回的 JSON 顶层必须是对象，且含非空字符串字段 document。",
        )

    # Legacy: plain-text full body only (non-JSON or non-object)
    return text.strip(), []


async def _run_document_review_llm(draft: str) -> tuple[str, list[dict[str, Any]]]:
    """Call the active chat model once; return (revised body, structured issues)."""
    from agentscope_runtime.engine.schemas.exception import AppBaseException

    from ...app.agent_context import get_current_agent_id
    from ...app.runner.title_generator import _consume_model_response
    from ..model_factory import create_model_and_formatter

    trimmed = (draft or "").strip()
    if len(trimmed) > _DOC_REVIEW_MAX_USER_CHARS:
        trimmed = trimmed[:_DOC_REVIEW_MAX_USER_CHARS].rstrip() + "\n\n[…正文已截断…]"

    agent_id = get_current_agent_id()
    try:
        model, _ = create_model_and_formatter(agent_id=agent_id)
    except (ValueError, AppBaseException, OSError) as exc:
        raise RuntimeError(
            "未配置可用的对话模型，无法生成修改后正文。",
        ) from exc

    messages = [
        {"role": "system", "content": _DOC_REVIEW_SYSTEM_ZH},
        {
            "role": "user",
            "content": (
                "以下为待审正文。请严格按系统说明只输出 JSON（含 document 与 "
                "resultList）：\n\n"
                f"{trimmed}"
            ),
        },
    ]
    try:
        revised = await asyncio.wait_for(
            _consume_model_response(model, messages),
            timeout=_DOC_REVIEW_LLM_TIMEOUT,
        )
    except asyncio.TimeoutError as exc:
        raise RuntimeError(
            f"文档审核超时（{_DOC_REVIEW_LLM_TIMEOUT:.0f}s）。",
        ) from exc

    return _parse_document_review_response(revised or "")


def _doc_reviewer_tool_locals_for_merge(
    content: str,
    file_path: str,
    path: str,
    start_line: int | None,
    end_line: int | None,
    extra: dict[str, Any],
) -> dict[str, Any]:
    """Build ``locals()``-like dict for :func:`_merge_body_from_tool_locals`."""
    loc: dict[str, Any] = {k: "" for k in _ALL_BODY_VALUE_KEYS}
    loc["content"] = content or ""
    for k in _GOV_INLINE_BODY_KEYS:
        if k == "content":
            continue
        val = extra.get(k)
        if isinstance(val, str):
            loc[k] = val
    data_v = extra.get("data")
    if isinstance(data_v, str):
        loc["data"] = data_v
    elif isinstance(data_v, dict):
        loc["data"] = json.dumps(data_v, ensure_ascii=False)
    else:
        loc["data"] = ""
    nr = extra.get("normalizedResult")
    if isinstance(nr, dict):
        doc = nr.get("document")
        if isinstance(doc, str) and doc.strip():
            loc["document"] = doc.strip()
    loc["file_path"] = file_path
    loc["path"] = path
    loc["start_line"] = start_line
    loc["end_line"] = end_line
    loc["extra"] = extra
    return loc


async def doc_reviewer(
    content: str = "",
    file_path: str = "",
    path: str = "",
    start_line: int | None = None,
    end_line: int | None = None,
    **extra: Any,
) -> ToolResponse:
    """Inline or file-based draft for review (校对、审阅、定稿).

    Exposed to the model as ``doc_reviewer``. Intended to review **the body
    produced by the previous ``gov_document_writer`` step** (plain ``content`` /
    ``document``, ``normalizedResult.document``, or a stringified / block-array
    tool JSON envelope). If that cannot be resolved, falls back to
    ``file_path`` / ``path`` (same resolution as :func:`read_file`).

    Returns a JSON envelope; ``normalizedResult.document`` is the
    **审核修改后正文** (LLM-revised full text). Top-level ``resultList`` holds
    structured change entries (``reason``, ``offsets``, ``errorType``, …).
    """
    ex = {k: v for k, v in extra.items()}
    prior_body = (
        _unwrap_gov_writer_prior_document(content)
        or _unwrap_gov_writer_prior_document(ex.get("data"))
        or _unwrap_gov_writer_prior_document(ex)
    )
    content_for_merge = prior_body if prior_body else (content or "")
    merge_loc = _doc_reviewer_tool_locals_for_merge(
        content_for_merge,
        file_path,
        path,
        start_line,
        end_line,
        ex,
    )
    merged = (_merge_body_from_tool_locals(merge_loc) or "").strip()
    max_bytes = get_current_recent_max_bytes() or DEFAULT_MAX_BYTES
    if merged:
        total_lines = merged.count("\n") + (1 if merged else 0)
        draft = truncate_text_output(
            merged,
            start_line=1,
            total_lines=total_lines,
            max_bytes=max_bytes,
            file_path="<inline-review>",
        )
    else:
        target = (file_path or path or "").strip()
        if not target:
            return _doc_reviewer_error(
                "错误：缺少待审阅内容。请先传入上一步的正文（content / document / "
                "normalizedResult 等），或提供 file_path / path（工作区文稿路径）。",
            )
        read_resp = await read_file(target, start_line=start_line, end_line=end_line)
        block0 = read_resp.content[0] if read_resp.content else None
        draft = _content_block_text(block0) or ""
        if draft.startswith("Error:"):
            return _doc_reviewer_error(draft)

    try:
        revised, result_list = await _run_document_review_llm(draft)
    except Exception as exc:  # noqa: BLE001
        logger.warning("doc_reviewer: revise LLM failed: %s", exc)
        return _doc_reviewer_error(str(exc))

    normalized_items = [
        _normalize_review_item(x) for x in result_list if isinstance(x, dict)
    ]

    out_lines = revised.count("\n") + (1 if revised else 0)
    document_out = truncate_text_output(
        revised,
        start_line=1,
        total_lines=out_lines,
        max_bytes=max_bytes,
        file_path="<revised-document>",
    )
    return _doc_reviewer_ok(document_out, normalized_items)


async def _save_gov_style_docx(
    *,
    skill_name_en: str,
    display_ok_zh: str,
    display_err_zh: str,
    content: str,
    title: str,
    file_name: str,
    document_type: str,
    machine_type: str,
    default_heading: str,
    file_path: str = "",
    path: str = "",
    start_line: int | None = None,
    end_line: int | None = None,
    save_path: str = "",
    savePath: str = "",
    include_normalized_result: bool = True,
    layout_result_list: dict[str, Any] | None = None,
    persist_docx: bool = True,
) -> ToolResponse:
    """Shared path: validate input, optionally write ``govdocs/*.docx``, JSON envelope."""
    body = (content or "").strip()
    if not body:
        target = (file_path or path or "").strip()
        if target:
            read_resp = await read_file(target, start_line=start_line, end_line=end_line)
            block0 = read_resp.content[0] if read_resp.content else None
            draft = _content_block_text(block0) or ""
            if draft.startswith("Error:"):
                return _gov_doc_tool_error(skill_name_en, display_err_zh, draft)
            body = draft.strip()
    if not body:
        return _gov_doc_tool_error(
            skill_name_en,
            display_err_zh,
            "错误：缺少正文。请提供：任一非空正文字段（如 content、text、document、"
            "body、markdown、keyword、data 等）、工作区 ``file_path`` / ``path`` 指向"
            "的文稿文件，或嵌套 JSON / 其它通过 ``**extra`` 传入的字段。",
        )

    heading = (title or "").strip()
    if not heading and (file_name or "").strip():
        heading = Path(file_name.strip()).stem
    if not heading:
        heading = default_heading

    doc_type_zh = (document_type or "").strip()
    mtype = (machine_type or "").strip()
    if doc_type_zh:
        type_label = doc_type_zh
    elif mtype and mtype.lower() != "notice":
        type_label = mtype
    else:
        type_label = ""

    document_plain = f"{heading}\n\n{body}".strip() + "\n"

    root: dict[str, Any] = {
        "skillName": skill_name_en,
        "stepIndex": _writer_step_index(),
        "displayText": display_ok_zh,
        "retryable": True,
        "sourceState": _SOURCE_OK,
        "errorDetail": None,
    }
    if persist_docx:
        out_hint = (savePath or save_path or "").strip()
        resolved = _resolve_gov_docx_output_path(out_hint, heading)
        resolved.parent.mkdir(parents=True, exist_ok=True)

        try:
            await asyncio.to_thread(
                _write_gov_docx,
                resolved,
                heading,
                body,
                type_label=type_label,
            )
        except Exception as e:
            return _gov_doc_tool_error(skill_name_en, display_err_zh, str(e))
        root["savePath"] = str(resolved.resolve())

    if include_normalized_result:
        root["normalizedResult"] = {
            "document": document_plain,
            "source": _SOURCE_OK,
        }
    if layout_result_list is not None:
        client_rl = _layout_result_list_for_client(layout_result_list)
        if client_rl is not None:
            root["resultList"] = client_rl
    return _gov_doc_tool_json(root)


async def gov_document_writer(
    content: str = "",
    title: str = "",
    file_name: str = "",
    document_type: str = "",
    type: str = "",  # noqa: A002  # API / model payload uses key "type"
    keyword: str = "",
    text: str = "",
    document: str = "",
    body: str = "",
    markdown: str = "",
    full_text: str = "",
    plain_text: str = "",
    article: str = "",
    source_text: str = "",
    message: str = "",
    draft: str = "",
    query: str = "",
    prompt: str = "",
    input: str = "",  # noqa: A002
    data: str = "",
    mainText: str = "",
    main_text: str = "",
    plainText: str = "",
    fullText: str = "",
    userContent: str = "",
    sourceText: str = "",
    file_path: str = "",
    path: str = "",
    start_line: int | None = None,
    end_line: int | None = None,
    save_path: str = "",
    savePath: str = "",
    **extra: Any,
) -> ToolResponse:
    """Draft a formal notice-style document (公文写作).

    Does **not** write ``.docx`` files or emit ``savePath``. Returns JSON with
    ``normalizedResult.document`` as plain ``标题 + 正文`` (no Markdown or machine
    ``type`` lines).

    Body text may arrive under many field names (see module constants), as JSON
    in ``data``, inside nested dicts, in ``**extra``, or only via ``file_path`` /
    ``path`` for a workspace file.
    """
    merged = _merge_body_from_tool_locals(locals())
    return await _save_gov_style_docx(
        skill_name_en=_SKILL_NAME_EN,
        display_ok_zh=_DISPLAY_OK_ZH,
        display_err_zh=_DISPLAY_ERR_ZH,
        content=merged,
        title=title,
        file_name=file_name,
        document_type=document_type,
        machine_type=type,
        default_heading="公文草稿",
        file_path=file_path,
        path=path,
        start_line=start_line,
        end_line=end_line,
        save_path=save_path,
        savePath=savePath,
        include_normalized_result=True,
        layout_result_list=None,
        persist_docx=False,
    )


def _layout_explicit_nonempty(rl: dict[str, Any]) -> bool:
    if not isinstance(rl, dict):
        return False
    t = rl.get("templates")
    if isinstance(t, list) and len(t) > 0:
        return True
    lst = rl.get("list")
    if isinstance(lst, list) and len(lst) > 0:
        return True
    tv = rl.get("total")
    try:
        return int(tv) > 0  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return False


def _template_title_for_layout_fetch(_locals: dict[str, Any]) -> str:
    """HaiRuo 模板检索关键字：``template_title`` 表示**公文标题**（非模板库内部名）。"""
    t = str(_locals.get("template_title") or "").strip()
    if t:
        return t
    extra_raw = _locals.get("extra")
    if isinstance(extra_raw, dict):
        v = extra_raw.get("template_title")
        if isinstance(v, str) and v.strip():
            return v.strip()
    return ""


async def gov_document_layout(
    content: str = "",
    title: str = "",
    file_name: str = "",
    document_type: str = "",
    type: str = "",  # noqa: A002
    keyword: str = "",
    text: str = "",
    document: str = "",
    body: str = "",
    markdown: str = "",
    full_text: str = "",
    plain_text: str = "",
    article: str = "",
    source_text: str = "",
    message: str = "",
    draft: str = "",
    query: str = "",
    prompt: str = "",
    input: str = "",  # noqa: A002
    data: str = "",
    mainText: str = "",
    main_text: str = "",
    plainText: str = "",
    fullText: str = "",
    userContent: str = "",
    sourceText: str = "",
    file_path: str = "",
    path: str = "",
    start_line: int | None = None,
    end_line: int | None = None,
    save_path: str = "",
    savePath: str = "",
    resultList: Any = None,
    templates: str = "",
    template_title: str = "",
    layout_page: int = 1,
    layout_page_size: int = 10,
    layout_defaults_path: str = "",
    **extra: Any,
) -> ToolResponse:
    """公文排版：仅按 ``template_title``（公文标题）推荐 HaiRuo 模板，**不写** ``.docx``。

    成功时返回非空的 ``savePath``：与原先写盘时相同的**待写入**绝对路径（本步不创建文件）。
    可传入显式 ``resultList`` / ``templates``（非空时跳过 HaiRuo 请求，仍须 ``template_title``）。
    不接受 ``templateTitle`` 驼峰键名。正文、``file_path`` 等字段会被忽略。
    """
    loc = {**locals()}
    tt = _template_title_for_layout_fetch(loc)
    if not tt:
        return _gov_doc_tool_error(
            _SKILL_NAME_LAYOUT_EN,
            _DISPLAY_ERR_LAYOUT_ZH,
            "错误：公文排版仅支持按公文标题推荐模板。请传入非空 ``template_title`` "
            "（公文标题）；本工具不再根据正文写盘。",
        )

    layout_rl_explicit = _layout_result_list_from_tool_locals(loc)
    fetched_rl: dict[str, Any] | None = None

    if _layout_explicit_nonempty(layout_rl_explicit):
        final_rl = layout_rl_explicit
    else:
        try:
            fetched_rl = await asyncio.to_thread(
                fetch_layout_template_result_list,
                tt,
                page=int(layout_page) if layout_page else 1,
                page_size=int(layout_page_size) if layout_page_size else 10,
                explicit_defaults_path=(layout_defaults_path or "").strip(),
            )
        except GovLayoutHttpError as exc:
            return _gov_doc_tool_error(
                _SKILL_NAME_LAYOUT_EN,
                _DISPLAY_ERR_LAYOUT_ZH,
                str(exc),
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning("gov_document_layout: layoutTemplate failed: %s", exc)
            return _gov_doc_tool_error(
                _SKILL_NAME_LAYOUT_EN,
                _DISPLAY_ERR_LAYOUT_ZH,
                str(exc),
            )
        final_rl = fetched_rl if fetched_rl is not None else layout_rl_explicit

    pending = _gov_layout_pending_save_absolute_path(
        template_title=tt,
        title=title,
        file_name=file_name,
        save_path=save_path,
        savePath=savePath,
    )
    return _gov_layout_ok_json_recommend_only(final_rl, pending_save_path=pending)
