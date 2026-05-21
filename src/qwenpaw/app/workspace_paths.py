# -*- coding: utf-8 -*-
"""Resolve and validate workspace file write targets."""
from __future__ import annotations

import base64
import io
import logging
import re
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import HTTPException

from ..constant import SECRET_DIR

GOVDOCS_SUBDIR = "govdocs"
GOVDOCS_SORT_ASC = "asc"
GOVDOCS_SORT_DESC = "desc"
DEFAULT_GOVDOCS_SORT_ORDER = GOVDOCS_SORT_DESC

# Max file size for embedding content in ``file_detail`` (10 MiB).
MAX_GOVDOCS_DETAIL_CONTENT_BYTES = 10 * 1024 * 1024

# Max binary payload for ``PUT /api/agentloop/workspace/files_binary`` (50 MiB).
MAX_WORKSPACE_BINARY_WRITE_BYTES = 50 * 1024 * 1024

# Max files and total bytes per ``GET /api/agentloop/workspace/download``.
MAX_WORKSPACE_DOWNLOAD_FILES = 100
MAX_WORKSPACE_DOWNLOAD_TOTAL_BYTES = 200 * 1024 * 1024

logger = logging.getLogger(__name__)

_DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_SUFFIX_MIME: dict[str, str] = {
    ".docx": _DOCX_MIME,
    ".doc": "application/msword",
}


def write_bytes_to_path(target: Path, data: bytes) -> None:
    """Create parent directories and write *data* to *target*."""
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)


def normalize_path_argument(path_str: str) -> str:
    """Normalize path query values (e.g. ``/C:/foo`` → ``C:/foo`` on Windows)."""
    normalized = (path_str or "").strip()
    if (
        len(normalized) >= 4
        and normalized[0] == "/"
        and normalized[2] == ":"
        and normalized[1].isalpha()
    ):
        normalized = normalized[1:]
    return normalized


def resolve_workspace_write_path(path_str: str, workspace_dir: Path) -> Path:
    """Resolve *path_str* to an absolute path under *workspace_dir*.

    Relative paths are resolved from *workspace_dir*. Absolute paths are allowed
    only when they resolve inside the workspace root.

    Raises:
        HTTPException: 400 for empty path, 403 for paths outside workspace or
            sensitive directories.
    """
    raw = normalize_path_argument(path_str)
    if not raw:
        raise HTTPException(status_code=400, detail="Path must not be empty")

    workspace_root = workspace_dir.expanduser().resolve()
    candidate = Path(raw).expanduser()
    if candidate.is_absolute():
        target = candidate.resolve()
    else:
        target = (workspace_root / candidate).resolve()

    try:
        target.relative_to(workspace_root)
    except ValueError as exc:
        raise HTTPException(
            status_code=403,
            detail="Path must be inside the agent workspace",
        ) from exc

    secret_root = SECRET_DIR.resolve()
    try:
        target.relative_to(secret_root)
        raise HTTPException(
            status_code=403,
            detail="Writing to the secrets directory is not allowed",
        )
    except ValueError:
        pass

    return target


def resolve_workspace_upload_directory(
    directory_str: str,
    workspace_dir: Path,
) -> Path:
    """Resolve *directory_str* to a directory path under *workspace_dir*."""
    target = resolve_workspace_write_path(directory_str, workspace_dir)
    if target.exists() and not target.is_dir():
        raise HTTPException(
            status_code=400,
            detail="directory must refer to a directory, not a file",
        )
    return target


def _validate_folder_name(folder_name: str) -> str:
    name = (folder_name or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="folder_name must not be empty")
    if name in (".", ".."):
        raise HTTPException(status_code=400, detail="folder_name is not allowed")
    if "/" in name or "\\" in name or Path(name).name != name:
        raise HTTPException(
            status_code=400,
            detail="folder_name must be a single folder name without path separators",
        )
    return name


def create_workspace_folder(
    parent_dir_str: str,
    folder_name: str,
    workspace_dir: Path,
) -> dict[str, Any]:
    """Create a folder named *folder_name* under *parent_dir_str*."""
    name = _validate_folder_name(folder_name)
    parent = resolve_workspace_upload_directory(parent_dir_str, workspace_dir)
    if not parent.exists():
        raise HTTPException(status_code=404, detail="parent_dir not found")
    if not parent.is_dir():
        raise HTTPException(status_code=400, detail="parent_dir must be a directory")

    target = resolve_unique_govdocs_dest(parent, name)

    try:
        target.mkdir(parents=False, exist_ok=False)
    except OSError as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Failed to create folder: {exc}",
        ) from exc

    govdocs_root = _govdocs_dir(workspace_dir)
    workspace_root = workspace_dir.expanduser().resolve()
    return _dir_entry(target, govdocs_root, workspace_root)


def _validate_file_name(file_name: str) -> str:
    raw = (file_name or "").strip()
    name = Path(raw).name
    if not name or name in {".", ".."}:
        raise HTTPException(status_code=400, detail="file_name must not be empty")
    if raw != name:
        raise HTTPException(
            status_code=400,
            detail="file_name must be a single file name without path separators",
        )
    return name


def _document_body_text(
    content_text: str | None,
    content_html: str | None,
) -> str:
    text = (content_text or "").strip()
    if text:
        return content_text or ""
    html = (content_html or "").strip()
    if not html:
        return ""
    plain = re.sub(r"<[^>]+>", "", html)
    plain = re.sub(r"\s+", " ", plain).strip()
    return plain


def _write_workspace_docx(path: Path, content_text: str | None, content_html: str | None) -> None:
    try:
        from docx import Document  # type: ignore[import-untyped]
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="python-docx is required to create .docx files",
        ) from exc

    doc = Document()
    body = _document_body_text(content_text, content_html)
    if body:
        for line in body.splitlines():
            doc.add_paragraph(line.rstrip("\r"))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def create_workspace_document(
    parent_dir_str: str,
    file_name: str,
    workspace_dir: Path,
    *,
    content_text: str | None = None,
    content_html: str | None = None,
    source: str | None = None,
) -> dict[str, Any]:
    """Create a file named *file_name* under *parent_dir_str*."""
    name = _validate_file_name(file_name)
    parent = resolve_workspace_upload_directory(parent_dir_str, workspace_dir)
    if not parent.exists():
        raise HTTPException(status_code=404, detail="parent_dir not found")
    if not parent.is_dir():
        raise HTTPException(status_code=400, detail="parent_dir must be a directory")

    target = resolve_unique_govdocs_dest(parent, name)

    suffix = target.suffix.lower()
    try:
        if suffix == ".docx":
            _write_workspace_docx(target, content_text, content_html)
        else:
            parent.mkdir(parents=True, exist_ok=True)
            payload = (content_text or "").encode("utf-8")
            write_bytes_to_path(target, payload)
    except HTTPException:
        raise
    except OSError as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Failed to create file: {exc}",
        ) from exc

    detail = govdocs_file_detail(target)
    if source:
        detail["source"] = source
    return detail


def save_upload_to_workspace_directory(
    directory: Path,
    upload_filename: str,
    data: bytes,
) -> dict[str, Any]:
    """Write uploaded bytes into *directory* using a safe, unique file name."""
    safe_name = Path((upload_filename or "").strip()).name
    if not safe_name or safe_name in {".", ".."}:
        raise HTTPException(status_code=400, detail="Upload filename must not be empty")

    directory.mkdir(parents=True, exist_ok=True)
    if not directory.is_dir():
        raise HTTPException(
            status_code=400,
            detail="directory must refer to a directory, not a file",
        )

    dest = resolve_unique_govdocs_dest(directory, safe_name)
    write_bytes_to_path(dest, data)
    return _file_entry(dest)


def mime_type_for_suffix(suffix: str) -> str:
    """Return Content-Type for a file suffix."""
    return _mime_for_suffix(suffix)


def resolve_workspace_download_file(
    path_str: str,
    workspace_dir: Path,
) -> Path:
    """Resolve *path_str* to an existing regular file under *workspace_dir*."""
    target = resolve_workspace_write_path(path_str, workspace_dir)
    if not target.is_file():
        raise HTTPException(status_code=404, detail=f"File not found: {path_str}")
    return target.resolve()


def resolve_workspace_download_files(
    path_strings: list[str],
    workspace_dir: Path,
) -> list[Path]:
    """Resolve and de-duplicate download targets; enforce count and size limits."""
    if not path_strings:
        raise HTTPException(status_code=400, detail="At least one path is required")
    if len(path_strings) > MAX_WORKSPACE_DOWNLOAD_FILES:
        raise HTTPException(
            status_code=400,
            detail=f"At most {MAX_WORKSPACE_DOWNLOAD_FILES} paths per request",
        )

    seen: set[str] = set()
    files: list[Path] = []
    total_bytes = 0

    for raw in path_strings:
        label = (raw or "").strip()
        if not label:
            raise HTTPException(status_code=400, detail="Path must not be empty")
        target = resolve_workspace_download_file(label, workspace_dir)
        key = str(target)
        if key in seen:
            continue
        seen.add(key)
        total_bytes += target.stat().st_size
        if total_bytes > MAX_WORKSPACE_DOWNLOAD_TOTAL_BYTES:
            limit_mb = MAX_WORKSPACE_DOWNLOAD_TOTAL_BYTES // (1024 * 1024)
            raise HTTPException(
                status_code=400,
                detail=f"Total download size exceeds {limit_mb} MB",
            )
        files.append(target)

    if not files:
        raise HTTPException(status_code=400, detail="No valid file paths provided")
    return files


def build_workspace_files_zip(
    workspace_dir: Path,
    files: list[Path],
) -> io.BytesIO:
    """Zip *files* using paths relative to *workspace_dir* as archive names."""
    workspace_root = workspace_dir.expanduser().resolve()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        used_arcnames: set[str] = set()
        for file_path in files:
            try:
                arcname = file_path.relative_to(workspace_root).as_posix()
            except ValueError:
                arcname = file_path.name
            if arcname in used_arcnames:
                parent = file_path.parent.name or "files"
                arcname = f"{parent}/{file_path.name}"
            used_arcnames.add(arcname)
            zf.write(file_path, arcname)
    buf.seek(0)
    return buf


def _govdocs_dir(workspace_dir: Path) -> Path:
    return workspace_dir.expanduser().resolve() / GOVDOCS_SUBDIR


def resolve_workspace_list_directory(
    path_str: str | None,
    workspace_dir: Path,
) -> Path:
    """Resolve the directory root for workspace ``get_list``.

    When *path_str* is omitted, defaults to ``govdocs/`` (may not exist yet).
    When provided, *path_str* must refer to an existing directory under the
    workspace.
    """
    govdocs_root = _govdocs_dir(workspace_dir)
    if not path_str or not path_str.strip():
        return govdocs_root

    target = resolve_workspace_write_path(path_str, workspace_dir)
    if not target.exists():
        raise HTTPException(status_code=404, detail="Path not found")
    if not target.is_dir():
        raise HTTPException(status_code=400, detail="Path must be a directory")
    return target


def list_govdocs_files(
    workspace_dir: Path,
    *,
    list_root: Path | None = None,
) -> list[dict[str, Any]]:
    """List files and directories under *list_root* recursively.

    Defaults to ``govdocs/``. Returns an empty list when the root directory
    does not exist (only for the default ``govdocs/`` case).
    """
    govdocs_root = _govdocs_dir(workspace_dir)
    workspace_root = workspace_dir.expanduser().resolve()
    root = list_root if list_root is not None else govdocs_root
    if not root.is_dir():
        return []

    entries: list[dict[str, Any]] = []
    for path in sorted(root.rglob("*")):
        if path.is_dir():
            entries.append(
                _dir_entry(path, govdocs_root, workspace_root),
            )
        elif path.is_file():
            entries.append(
                _file_entry(path, govdocs_root=govdocs_root, workspace_root=workspace_root),
            )
    return entries


def filter_govdocs_by_filename(
    items: list[dict[str, Any]],
    filename_query: str | None,
) -> list[dict[str, Any]]:
    """Keep items matching *filename_query* against name or ``relative_path``."""
    if not filename_query or not filename_query.strip():
        return items
    needle = filename_query.strip().casefold()
    return [item for item in items if _govdocs_entry_matches_query(item, needle)]


def _govdocs_entry_matches_query(item: dict[str, Any], needle: str) -> bool:
    if needle in item.get("filename", "").casefold():
        return True
    relative = item.get("relative_path", "")
    return bool(relative) and needle in relative.casefold()


def sort_govdocs_by_modified_time(
    items: list[dict[str, Any]],
    sort_order: str = DEFAULT_GOVDOCS_SORT_ORDER,
) -> list[dict[str, Any]]:
    """Sort by ``modified_time`` ascending or descending."""
    order = (sort_order or DEFAULT_GOVDOCS_SORT_ORDER).strip().lower()
    if order not in (GOVDOCS_SORT_ASC, GOVDOCS_SORT_DESC):
        raise HTTPException(
            status_code=400,
            detail=f"sortOrder must be '{GOVDOCS_SORT_ASC}' or '{GOVDOCS_SORT_DESC}'",
        )
    reverse = order == GOVDOCS_SORT_DESC
    return sorted(items, key=lambda item: item["modified_time"], reverse=reverse)


def annotate_govdocs_user(
    items: list[dict[str, Any]],
    user: str,
) -> list[dict[str, Any]]:
    """Add ``user`` to each list entry."""
    return [{**item, "user": user} for item in items]


def build_govdocs_list_response(items: list[dict[str, Any]]) -> dict[str, Any]:
    """Wrap govdocs entries for ``get_list`` (full list, no pagination)."""
    return {"list": items}


def resolve_govdocs_file_path(path_str: str, workspace_dir: Path) -> Path:
    """Resolve *path_str* to a file path that must lie under ``govdocs/``."""
    target = resolve_workspace_write_path(path_str, workspace_dir)
    govdocs_root = _govdocs_dir(workspace_dir).resolve()
    try:
        target.relative_to(govdocs_root)
    except ValueError as exc:
        raise HTTPException(
            status_code=403,
            detail="Path must be under the govdocs directory",
        ) from exc
    return target


def _entry_relative_path(
    path: Path,
    *,
    govdocs_root: Path,
    workspace_root: Path,
) -> str:
    resolved = path.resolve()
    try:
        return resolved.relative_to(govdocs_root.resolve()).as_posix()
    except ValueError:
        return resolved.relative_to(workspace_root).as_posix()


def _dir_entry(
    path: Path,
    govdocs_root: Path,
    workspace_root: Path,
) -> dict[str, Any]:
    stat = path.stat()
    return {
        "type": "directory",
        "filename": path.name,
        "relative_path": _entry_relative_path(
            path,
            govdocs_root=govdocs_root,
            workspace_root=workspace_root,
        ),
        "path": str(path.resolve()),
        "size": 0,
        "suffix": "",
        "created_time": datetime.fromtimestamp(
            stat.st_ctime,
            tz=timezone.utc,
        ).isoformat(),
        "modified_time": datetime.fromtimestamp(
            stat.st_mtime,
            tz=timezone.utc,
        ).isoformat(),
    }


def _file_entry(
    path: Path,
    *,
    govdocs_root: Path | None = None,
    workspace_root: Path | None = None,
) -> dict[str, Any]:
    stat = path.stat()
    entry: dict[str, Any] = {
        "type": "file",
        "filename": path.name,
        "path": str(path.resolve()),
        "size": stat.st_size,
        "suffix": path.suffix.lower(),
        "created_time": datetime.fromtimestamp(
            stat.st_ctime,
            tz=timezone.utc,
        ).isoformat(),
        "modified_time": datetime.fromtimestamp(
            stat.st_mtime,
            tz=timezone.utc,
        ).isoformat(),
    }
    if govdocs_root is not None and workspace_root is not None:
        entry["relative_path"] = _entry_relative_path(
            path,
            govdocs_root=govdocs_root,
            workspace_root=workspace_root,
        )
    return entry


def _mime_for_suffix(suffix: str) -> str:
    return _SUFFIX_MIME.get(suffix.lower(), "application/octet-stream")


def _extract_docx_plain_text(path: Path) -> str | None:
    """Extract paragraph text from a ``.docx`` file (requires ``python-docx``)."""
    try:
        from docx import Document  # type: ignore[import-untyped]
    except ImportError:
        logger.warning("python-docx not installed; contentText omitted for %s", path)
        return None

    try:
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        return "\n".join(parts).strip()
    except Exception as exc:  # noqa: BLE001
        logger.warning("Failed to extract docx text from %s: %s", path, exc)
        return None


def govdocs_file_detail(path: Path) -> dict[str, Any]:
    """Return metadata and file content for an existing ``govdocs`` file."""
    if not path.is_file():
        raise HTTPException(status_code=404, detail="File not found")

    detail = _file_entry(path)
    size = int(detail["size"])
    detail["contentType"] = _mime_for_suffix(path.suffix)

    if size > MAX_GOVDOCS_DETAIL_CONTENT_BYTES:
        detail["contentOmitted"] = True
        detail["contentOmittedReason"] = (
            f"File exceeds {MAX_GOVDOCS_DETAIL_CONTENT_BYTES // (1024 * 1024)} MB; "
            "content not included in response"
        )
        detail["contentBase64"] = None
        detail["contentText"] = None
        return detail

    raw = path.read_bytes()
    detail["contentBase64"] = (
        base64.b64encode(raw).decode("ascii") if raw else ""
    )
    detail["contentOmitted"] = False

    if path.suffix.lower() == ".docx":
        detail["contentText"] = _extract_docx_plain_text(path)
    else:
        detail["contentText"] = None

    return detail


_GOVDOCS_COPY_LABEL = " - 副本"
_MAX_GOVDOCS_COPY_SUFFIX_ATTEMPTS = 999


def resolve_unique_govdocs_dest(
    parent: Path,
    desired_filename: str,
    *,
    exclude: Path | None = None,
) -> Path:
    """Return a destination path for *desired_filename*, avoiding name clashes.

    If the desired name is taken, append `` - 副本``, then `` - 副本 (2)``, etc.
    """
    safe_name = Path(desired_filename.strip()).name
    if not safe_name or safe_name in {".", ".."}:
        raise HTTPException(status_code=400, detail="newFilename must not be empty")

    def _available(candidate: Path) -> bool:
        if not candidate.exists():
            return True
        if exclude is not None and candidate.resolve() == exclude.resolve():
            return True
        return False

    first = parent / safe_name
    if _available(first):
        return first

    stem = Path(safe_name).stem
    suffix = Path(safe_name).suffix

    copy_first = parent / f"{stem}{_GOVDOCS_COPY_LABEL}{suffix}"
    if _available(copy_first):
        return copy_first

    for n in range(2, _MAX_GOVDOCS_COPY_SUFFIX_ATTEMPTS + 1):
        candidate = parent / f"{stem}{_GOVDOCS_COPY_LABEL} ({n}){suffix}"
        if _available(candidate):
            return candidate

    raise HTTPException(
        status_code=500,
        detail="Could not allocate a unique file name",
    )


def _assert_not_govdocs_root(path: Path, workspace_dir: Path) -> None:
    if path.resolve() == _govdocs_dir(workspace_dir).resolve():
        raise HTTPException(
            status_code=400,
            detail="Cannot modify the govdocs root directory",
        )


def govdocs_path_detail(path: Path, workspace_dir: Path) -> dict[str, Any]:
    """Return metadata for a file or directory under ``govdocs/``."""
    if path.is_file():
        return govdocs_file_detail(path)
    if path.is_dir():
        govdocs_root = _govdocs_dir(workspace_dir)
        workspace_root = workspace_dir.expanduser().resolve()
        return _dir_entry(path, govdocs_root, workspace_root)
    raise HTTPException(status_code=404, detail="Path not found")


def rename_govdocs_file(
    path: Path,
    new_filename: str,
    workspace_dir: Path,
) -> Path:
    """Rename a file or directory to *new_filename* (basename only) in ``govdocs/``."""
    if not path.exists():
        raise HTTPException(status_code=404, detail="Path not found")
    _assert_not_govdocs_root(path, workspace_dir)

    safe_name = _validate_file_name(new_filename)
    dest = resolve_unique_govdocs_dest(
        path.parent,
        safe_name,
        exclude=path.resolve(),
    )
    path.rename(dest)
    return dest.resolve()


def delete_govdocs_file(path: Path, workspace_dir: Path) -> None:
    """Delete a file or directory under ``govdocs/`` (recursive for directories)."""
    if not path.exists():
        raise HTTPException(status_code=404, detail="Path not found")
    _assert_not_govdocs_root(path, workspace_dir)

    if path.is_file():
        path.unlink()
    elif path.is_dir():
        shutil.rmtree(path)
    else:
        raise HTTPException(status_code=404, detail="Path not found")


MAX_GOVDOCS_BATCH_DELETE = 100


def delete_govdocs_files_batch(
    workspace_dir: Path,
    path_strings: list[str],
) -> dict[str, Any]:
    """Delete multiple files or directories under ``govdocs/``; continue on failures."""
    if not path_strings:
        raise HTTPException(status_code=400, detail="paths must not be empty")
    if len(path_strings) > MAX_GOVDOCS_BATCH_DELETE:
        raise HTTPException(
            status_code=400,
            detail=f"At most {MAX_GOVDOCS_BATCH_DELETE} paths per request",
        )

    deleted: list[dict[str, str]] = []
    failed: list[dict[str, str]] = []

    for raw in path_strings:
        path_label = (raw or "").strip()
        if not path_label:
            failed.append({"path": raw or "", "detail": "Path must not be empty"})
            continue
        try:
            target = resolve_govdocs_file_path(path_label, workspace_dir)
            delete_govdocs_file(target, workspace_dir)
            deleted.append({"path": str(target.resolve())})
        except HTTPException as exc:
            detail = exc.detail
            if not isinstance(detail, str):
                detail = str(detail)
            failed.append({"path": path_label, "detail": detail})
        except OSError as exc:
            failed.append({"path": path_label, "detail": str(exc)})

    return {
        "deleted": deleted,
        "failed": failed,
        "deletedCount": len(deleted),
        "failedCount": len(failed),
    }
